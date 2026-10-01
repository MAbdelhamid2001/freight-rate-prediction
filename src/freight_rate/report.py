"""Model report: train vs test accuracy for every model family, plus the
submission report (DOCX) and its Markdown twin for the repository.

    uv run freight-report            # after freight-train, freight-predict and score.py

For the best configuration of each model family (taken from the MLflow study
in reports/experiments.csv) it refits on each time fold and records accuracy
on the rows it trained on and on the two months it did not see. The gap
between the two is the overfitting each family shows.

Writes
  reports/model_report.csv          one row per family x fold x split
  reports/model_summary.csv         one row per family (means over folds)
  reports/freight_rate_report.docx  the submission report
  reports/REPORT.md                 the same content, readable on GitHub
"""

from __future__ import annotations

import argparse
import json
import logging
import math
from pathlib import Path

import numpy as np
import pandas as pd

from freight_rate import config
from freight_rate.data import load_training_data, load_validation_data
from freight_rate.features import ReferenceData
from freight_rate.train import Candidate, cached_noise_flags, fit_candidate
from freight_rate.validation import regression_metrics, time_folds

log = logging.getLogger("freight_rate.report")

MODEL_REPORT_PATH = config.REPORTS_DIR / "model_report.csv"
MODEL_SUMMARY_PATH = config.REPORTS_DIR / "model_summary.csv"
DOCX_PATH = config.REPORTS_DIR / "freight_rate_report.docx"
MARKDOWN_PATH = config.REPORTS_DIR / "REPORT.md"
SCORER_CHART = config.PROJECT_ROOT / "scorer_results" / "candidate_december.png"
COMMITTED_CHART = config.FIGURES_DIR / "10_candidate_december.png"

INT_PARAMS = {"width", "depth", "batch_size", "n_estimators", "num_leaves", "min_child_samples", "max_depth", "n_knots", "max_leaf_nodes",
              "min_samples_leaf", "max_iter", "gam_n_knots", "gbm_num_leaves", "gbm_min_child_samples",
              "gbm_n_estimators"}
FAMILY_LABELS = {
    "baseline_rpm": "Baseline ($/mile by equipment x distance band)",
    "ridge": "Spline GAM (ColumnTransformer: impute, scale, B-splines -> Ridge)",
    "random_forest": "Random Forest",
    "hist_gbm": "HistGradientBoosting",
    "xgboost": "XGBoost",
    "lightgbm": "LightGBM",
    "blend": "Blend (GAM + LightGBM, log space)",
    "mlp": "Neural network (MLP, scaled inputs)",
}


# ---------------------------------------------------------------------------
# Train vs test accuracy per family
# ---------------------------------------------------------------------------
def _params_from_row(row: pd.Series) -> dict:
    if isinstance(row.get("params_json"), str):
        return json.loads(row["params_json"])
    params = {}
    for col, value in row.items():
        if not col.startswith("hp_") or value is None or (isinstance(value, float) and math.isnan(value)):
            continue
        name = col[3:]
        if name in INT_PARAMS and float(value).is_integer():
            value = int(value)
        params[name] = value
    # A tuned forest whose max_depth is blank was searched with max_depth=None.
    if row["model"] == "random_forest" and params and "max_depth" not in params:
        params["max_depth"] = None
    return params


def best_candidates(experiments: pd.DataFrame) -> list[Candidate]:
    pool = experiments[experiments.stage.isin(["features", "recency", "models", "tuning"])]
    best = pool.sort_values(config.SELECTION_METRIC).drop_duplicates("model")
    candidates = []
    for _, row in best.sort_values(config.SELECTION_METRIC).iterrows():
        half_life = row.get("half_life_days")
        candidates.append(Candidate(
            model=row["model"], feature_set=row["feature_set"],
            half_life_days=None if pd.isna(half_life) else float(half_life),
            params=_params_from_row(row),
        ))
    return candidates


def train_vs_test(candidates: list[Candidate], frame: pd.DataFrame, reference: ReferenceData,
                  noisy: np.ndarray) -> pd.DataFrame:
    y = frame[config.TARGET].to_numpy(dtype=float)
    rows = []
    for cand in candidates:
        log.info("train vs test: %s", cand.name)
        for fold in time_folds(frame):
            train_idx = fold.train_idx[~noisy[fold.train_idx]]
            model = fit_candidate(cand, frame.iloc[train_idx], reference)
            test_pred = model.predict(frame.iloc[fold.test_idx])
            clean_test = ~noisy[fold.test_idx]
            splits = {
                "train (in-sample)": (train_idx, model.predict(frame.iloc[train_idx])),
                "test (next 2 months)": (fold.test_idx[clean_test], test_pred[clean_test]),
                "test incl. corrupted labels": (fold.test_idx, test_pred),
            }
            for split, (idx, pred) in splits.items():
                metrics = regression_metrics(y[idx], pred)
                within5 = float(np.mean(np.abs(pred - y[idx]) / y[idx] < 0.05) * 100)
                rows.append({"model": cand.model, "candidate": cand.name, "fold": fold.name, "split": split,
                             "rows": len(idx), **metrics, "within_5pct": within5})
    return pd.DataFrame(rows)


def summarise(report: pd.DataFrame) -> pd.DataFrame:
    wide = report.pivot_table(index=["model", "candidate"], columns="split",
                              values=["mape", "mae", "r2", "within_5pct"], aggfunc="mean")
    out = pd.DataFrame({
        "model": [i[0] for i in wide.index],
        "configuration": [i[1] for i in wide.index],
        "train_mape": wide[("mape", "train (in-sample)")].to_numpy(),
        "test_mape": wide[("mape", "test (next 2 months)")].to_numpy(),
        "train_mae": wide[("mae", "train (in-sample)")].to_numpy(),
        "test_mae": wide[("mae", "test (next 2 months)")].to_numpy(),
        "test_mae_incl_corrupted": wide[("mae", "test incl. corrupted labels")].to_numpy(),
        "train_r2": wide[("r2", "train (in-sample)")].to_numpy(),
        "test_r2": wide[("r2", "test (next 2 months)")].to_numpy(),
        "test_within_5pct": wide[("within_5pct", "test (next 2 months)")].to_numpy(),
    })
    out["overfit_gap_pts"] = out.test_mape - out.train_mape
    return out.sort_values("test_mape").reset_index(drop=True)


# ---------------------------------------------------------------------------
# Document content (rendered to DOCX and Markdown from the same blocks)
# ---------------------------------------------------------------------------
def _fmt_table(df: pd.DataFrame, decimals: int = 2) -> pd.DataFrame:
    out = df.copy()
    for col in out.columns:
        if pd.api.types.is_float_dtype(out[col]):
            out[col] = out[col].map(lambda v: "" if pd.isna(v) else f"{v:,.{decimals}f}")
    return out.astype(str)


def build_blocks(summary: pd.DataFrame) -> list[tuple]:
    card = json.loads(config.MODEL_CARD_PATH.read_text(encoding="utf-8"))
    evaluation = json.loads(config.EVALUATION_PATH.read_text(encoding="utf-8"))
    eda = json.loads(config.EDA_SUMMARY_PATH.read_text(encoding="utf-8"))
    experiments = pd.read_csv(config.EXPERIMENTS_PATH)
    checks = card["robustness_checks"]
    holdout = evaluation["holdout"]
    q = eda["quality"]["train"]
    cv = card["cv_metrics"]
    winner = summary[summary.model == card["estimator"]].iloc[0]
    baseline = summary[summary.model == "baseline_rpm"]
    fig = config.FIGURES_DIR

    feature_table = (experiments[experiments.stage == "features"]
                     .pivot_table(index="feature_set", columns="model", values=config.SELECTION_METRIC)
                     .reset_index().rename(columns={"feature_set": "feature set",
                                                    "lightgbm": "LightGBM CV MAPE %",
                                                    "ridge": "Spline GAM CV MAPE %"}))
    model_table = summary.assign(model=summary.model.map(FAMILY_LABELS).fillna(summary.model))[
        ["model", "train_mape", "test_mape", "overfit_gap_pts", "test_mae", "test_mae_incl_corrupted",
         "test_r2", "test_within_5pct"]].rename(columns={
            "train_mape": "Train MAPE %", "test_mape": "Test MAPE %", "overfit_gap_pts": "Gap (pts)",
            "test_mae": "Test MAE $", "test_mae_incl_corrupted": "Test MAE $ (all labels)",
            "test_r2": "Test R²", "test_within_5pct": "Test within ±5 %"})
    folds = pd.DataFrame(card["cv_per_fold"])[["fold", "n_train", "n_test", "clean_mape", "clean_mae",
                                                "clean_r2", "clean_bias", "all_mape", "all_mae"]]
    folds.columns = ["Test months", "Train rows", "Test rows", "MAPE %", "MAE $", "R²", "Bias $",
                     "MAPE % (all labels)", "MAE $ (all labels)"]
    segments = pd.DataFrame(holdout["by_equipment"] + holdout["by_distance_band"])[
        ["segment", "loads", "mape", "mae", "bias"]].rename(columns={
            "segment": "Segment", "loads": "Loads", "mape": "MAPE %", "mae": "MAE $", "bias": "Bias $"})
    tuned = experiments[experiments.stage == "tuning"].groupby("model")[config.SELECTION_METRIC].agg(
        ["count", "min"]).reset_index()
    defaults = experiments[experiments.stage.isin(["models", "recency", "features"])].groupby("model")[
        config.SELECTION_METRIC].min()
    tuned["default"] = tuned.model.map(defaults)
    tuned = tuned.rename(columns={"model": "Family", "count": "Optuna trials", "min": "Best tuned CV MAPE %",
                                  "default": "Default CV MAPE %"})
    inference = evaluation["inference"]
    sanity = evaluation["validation_predictions"]
    sanity_table = pd.DataFrame(sanity["by_equipment"]).T.reset_index().rename(columns={
        "index": "Equipment", "predicted_median_rpm": "Nov-Dec predicted median $/mile",
        "sep_oct_median_rpm": "Sep-Oct actual median $/mile"})
    # The committed copy under reports/figures, so the Markdown link works on GitHub.
    december_chart = COMMITTED_CHART if COMMITTED_CHART.is_file() else SCORER_CHART

    blocks: list[tuple] = [
        ("title", "Freight Rate Prediction: Approach, Validation and Results"),
        ("p", f"Model: {FAMILY_LABELS.get(card['estimator'], card['estimator'])}, feature set "
              f"'{card['feature_set']}', trained on {card['training_rows']:,} loads "
              f"({card['training_period'][0]} to {card['training_period'][1]}). "
              f"Selected from {len(experiments)} MLflow-tracked runs."),
        ("h1", "1. Summary"),
        ("bullets", [
            f"Out-of-time accuracy (mean of 3 folds, each predicting the next two months): MAPE "
            f"{cv['cv_clean_mape']:.2f}%, MAE ${cv['cv_clean_mae']:.0f} on correctly labelled loads; "
            f"{winner.test_within_5pct:.1f}% of loads priced within ±5%.",
            f"Latest fold (train Jan-Aug, predict Sep-Oct, the closest rehearsal of the real task): "
            f"MAPE {holdout['overall_clean']['mape']:.2f}%, R² {holdout['overall_clean']['r2']:.4f}.",
            (f"The $/mile rule of thumb scores {baseline.test_mape.iloc[0]:.2f}% on the same folds, so the "
             f"model cuts its error by {(1 - winner.test_mape / baseline.test_mape.iloc[0]) * 100:.0f}%.")
            if not baseline.empty else "",
            f"About {card['dropped_noisy_labels']} labels "
            f"({card['dropped_noisy_labels'] / (card['dropped_noisy_labels'] + card['training_rows']) * 100:.1f}%) are "
            "corrupted (2-5x off). Metrics are reported with and without them; no model can predict them.",
        ]),
        ("h1", "2. Data exploration: key findings"),
        ("bullets", [
            f"{eda['rows']['train']:,} labelled loads (Jan-Oct 2025) and {eda['rows']['validation']:,} loads "
            "to price (Nov-Dec 2025): the task is a two-month forecast, not interpolation.",
            "Distance explains most of the rate (correlation 0.91). $/mile falls with distance and is "
            "highest for Reefer, then Flatbed, then Dry Van.",
            f"market_index is one value per day plus small noise (within-day std "
            f"{eda['market_shift']['within_day_std']:.3f} vs {eda['market_shift']['across_day_std']:.3f} across "
            "days). It has a weekly cycle on top of a piecewise trend. Rates follow the trend loosely, but the "
            "mapping drifts: Sep-Oct rates are ~5% above Jan-Feb at the same market level.",
            "quote_signal looks irrelevant by correlation (0.10) but has a V-shaped effect whose direction "
            "flips from month to month, so it is unstable as a forward-looking signal.",
            f"{eda['unseen_city_rows_in_validation']:,} validation loads (12%) touch one of 8 cities that never "
            "appear in training (e.g. Chicago, Charlotte, Laredo).",
        ]),
        ("img", fig / "02_distance_equipment.png", "Rate vs distance and $/mile by equipment"),
        ("img", fig / "03_market_over_time.png", "Market index and daily $/mile over time"),
        ("img", fig / "06_signals.png", "quote_signal effect and correlations"),
        ("h1", "3. Data-quality issues and how they were handled"),
        ("table", pd.DataFrame([
            ["Corrupted posted rates", f"~{card['dropped_noisy_labels']} rows (1.4%)",
             "Flagged out-of-fold (label vs. prediction from a model that never saw the row, >1.6x off); "
             "dropped from training, excluded from 'clean' metrics."],
            ["Negative weight", f"{q['negative_weight']}", "Sign error (same distribution as positives): abs()."],
            ["Missing weight", f"{q['missing_weight']}", "Median weight of the same equipment type, fitted in the pipeline."],
            ["Missing market_index", f"{q['missing_market_index']}",
             "That day's mean market index (error ~0.025 instead of ~0.17 for a global median)."],
            ["Weight at 47,500 lb cap", f"{q['weight_at_cap']}", "Kept: a legal-maximum cap, not an error."],
            ["Distance/straight-line > 1.6", f"{q['circuity_outliers']}",
             "Road distance kept; the circuity feature is clipped to [0.8, 2.5]."],
            ["Unseen cities (validation)", f"{eda['unseen_city_rows_in_validation']}",
             "No city IDs: coordinates, distance and bearing generalise to new cities."],
        ], columns=["Issue", "Rows", "Treatment"])),
        ("img", fig / "04_data_quality.png", "Weight and circuity defects"),
        ("h1", "4. Validation approach and data split"),
        ("p", "Labels end on 31 Oct and the loads to price are in Nov-Dec, so validation must answer one "
              "question: how well does a model trained on the past price the next two months? The development "
              "data is split by time into three expanding-window folds with a two-month test block, which is "
              "the same shape as the real task:"),
        ("table", pd.DataFrame(
            [[f"Jan-{(pd.Timestamp(train_end) - pd.Timedelta(days=1)):%b}",
              f"{pd.Timestamp(test_start):%b}-{(pd.Timestamp(test_end) - pd.Timedelta(days=1)):%b}"]
             for train_end, test_start, test_end in config.CV_FOLDS],
            columns=["Train on", "Test on"])),
        ("bullets", [
            "Selection metric: mean MAPE over the three folds on correctly labelled loads. MAPE because a $40 "
            "miss matters on a $300 load and not on a $6,000 one.",
            "No random split: the market and quote relationships drift month to month, and only a forward "
            f"split measures that. (A shuffled 3-fold on the chosen model scores "
            f"{checks['random_split_cv_clean_mape']:.2f}% vs {checks['time_split_cv_clean_mape']:.2f}% forward. "
            "The model has no calendar features, so shuffling mixes every month's level shift into each test "
            "fold rather than leaking.)",
            "Imputation medians and recency weights are fitted inside each fold's training window. "
            "Label-noise flags are computed once, out-of-fold over all labelled loads (each load is judged by a "
            "model that never saw it), and applied identically to every candidate. The market calendar and "
            "city coordinates use feature columns only, never the target.",
            f"Unseen-city rehearsal: hiding 8 cities from training gives {checks['unseen_city_clean_mape']:.2f}% "
            f"on loads touching them vs {checks['seen_city_clean_mape']:.2f}% on the rest.",
            f"Label-noise filter: training on all labels instead gives {checks['no_noise_filter_cv_clean_mape']:.2f}% "
            f"(vs {checks['time_split_cv_clean_mape']:.2f}%).",
        ]),
        ("h1", "5. Experiments (MLflow)"),
        ("p", "Stage 1 compared feature sets under a tree model and an additive model. Every market-index "
              "variant and the trend feature lose forward in time, and dropping quote_signal helps both "
              "families:"),
        ("table", _fmt_table(feature_table, 3)),
        ("p", "Stage 2 tested recency weighting (half-life 30-240 days). Stage 3 ran the model zoo with defaults. "
              "Stage 4 ran an Optuna search per family, every trial a nested MLflow run:"),
        ("table", _fmt_table(tuned, 3)),
        ("h1", "6. Models: accuracy on training vs test data"),
        ("p", "Best configuration per family, refitted on each fold. Train = the rows it was fitted on; "
              "Test = the following two months (clean labels). Gap = test minus train MAPE:"),
        ("table", _fmt_table(model_table, 2)),
        ("p", "Most models score better on the unseen months than on their own training rows. That is not "
              "leakage: the training window contains whole months whose rate level was shifted (January about "
              "5% cheap, June about 4.5% expensive), and the models deliberately have no calendar features, so "
              "they cannot memorise those month-specific shocks. That error appears in-sample, while the test "
              "months sit near the average level. Random Forest is the exception: it fits its training rows "
              "far more closely (fully grown trees memorise noise) and is the family that generalises worst. "
              "A negative or small gap is the behaviour we want from a model that must price future months."),
        ("img", fig / "08_model_comparison.png", "Best forward-CV MAPE per model family"),
        ("p", f"Why {FAMILY_LABELS.get(card['estimator'], card['estimator'])}: it has the lowest error on unseen "
              "months, the smallest train-test gap, and the most stable error across folds. The tree ensembles "
              "fit the training months more closely but react to every month-specific shift, which does not "
              "carry forward. The additive spline model captures the dominant, smooth effects (distance taper, "
              "equipment premium, weight, geography) and extrapolates them predictably."),
        ("h1", "7. Selected model: evaluation"),
        ("table", _fmt_table(folds, 2)),
        ("p", f"Holdout {holdout['fold']}: {holdout['within_5pct_clean']:.1f}% of loads within ±5% and "
              f"{holdout['within_10pct_clean']:.1f}% within ±10%. Error by segment:"),
        ("table", _fmt_table(segments, 2)),
        ("img", fig / "07_holdout_diagnostics.png", "Holdout diagnostics"),
        ("h1", "8. Validation set (Nov-Dec) predictions and inference"),
        ("p", f"All {sanity['rows']:,} loads in validation.csv were priced (${sanity['min']:,.0f} to "
              f"${sanity['max']:,.0f}, median ${sanity['median']:,.0f}, no non-positive values) and written to "
              "validation_predictions.csv in the template's order. The file has no labels, so its accuracy is "
              "measured by the assessor; the folds above are the estimate. Predicted $/mile lines up with the "
              "most recent actuals:"),
        ("table", _fmt_table(sanity_table, 3)),
        ("p", f"Inference: {inference['single_load_ms_p50']:.1f} ms per single load (p50, full pipeline incl. "
              f"cleaning and features), {inference['batch_rows_per_second']:,} loads/s in batch."),
        ("h1", "9. December chart (score.py)"),
        ("p", "Fixed load Lexington to Fort Wayne, 360 miles, Dry Van, 32,000 lb; only the date changes. "
              "The chart input has no coordinates or market columns, so coordinates come from the city "
              "lookup learned from the data. The model's date sensitivity is the learned day-of-week pattern; "
              "it has no month or trend term to extrapolate, by design."),
        ("img", december_chart, "Generated by the provided score.py"),
        ("h1", "10. Reproducing"),
        ("code", "uv sync\nuv run freight-eda\nuv run freight-train\nuv run freight-predict\n"
                 "python score.py --predictions validation_predictions.csv "
                 "--december-predictions data/december_chart_inputs.csv\nuv run freight-report"),
    ]
    return [b for b in blocks if not (b[0] == "bullets" and not any(b[1]))]


def render_docx(blocks: list[tuple], path: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.8)

    for block in blocks:
        kind = block[0]
        if kind == "title":
            doc.add_heading(block[1], level=0)
        elif kind == "h1":
            doc.add_heading(block[1], level=1)
        elif kind == "p":
            doc.add_paragraph(block[1])
        elif kind == "bullets":
            for item in block[1]:
                if item:
                    doc.add_paragraph(item, style="List Bullet")
        elif kind == "code":
            para = doc.add_paragraph()
            run = para.add_run(block[1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif kind == "table":
            df = block[1]
            table = doc.add_table(rows=1, cols=len(df.columns))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, name in zip(table.rows[0].cells, df.columns):
                cell.text = str(name)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.size = Pt(9)
            for _, row in df.iterrows():
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = str(value)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9)
            doc.add_paragraph()
        elif kind == "img":
            if Path(block[1]).is_file():
                doc.add_picture(str(block[1]), width=Inches(6.6))
                caption = doc.add_paragraph(block[2])
                caption.runs[0].font.size = Pt(8.5)
                caption.runs[0].font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
    doc.save(path)


def render_markdown(blocks: list[tuple], path: Path) -> None:
    lines: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind == "title":
            lines += [f"# {block[1]}", ""]
        elif kind == "h1":
            lines += [f"## {block[1]}", ""]
        elif kind == "p":
            lines += [block[1], ""]
        elif kind == "bullets":
            lines += [f"- {item}" for item in block[1] if item] + [""]
        elif kind == "code":
            lines += ["```bash", block[1], "```", ""]
        elif kind == "table":
            df = block[1]
            lines.append("| " + " | ".join(map(str, df.columns)) + " |")
            lines.append("|" + "---|" * len(df.columns))
            lines += ["| " + " | ".join(map(str, row)) + " |" for row in df.itertuples(index=False)]
            lines.append("")
        elif kind == "img":
            if Path(block[1]).is_file():
                rel = Path(block[1]).resolve().relative_to(config.REPORTS_DIR.resolve()) \
                    if config.REPORTS_DIR.resolve() in Path(block[1]).resolve().parents else Path(block[1]).name
                lines += [f"![{block[2]}]({Path(rel).as_posix()})", f"*{block[2]}*", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Build the model report and the submission document.")
    parser.add_argument("--skip-refit", action="store_true",
                        help="reuse reports/model_report.csv instead of refitting every family")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")

    if args.skip_refit and MODEL_REPORT_PATH.is_file():
        report = pd.read_csv(MODEL_REPORT_PATH)
    else:
        train = load_training_data()
        reference = ReferenceData.from_frames(train, load_validation_data())
        noisy = cached_noise_flags(train, reference)
        candidates = best_candidates(pd.read_csv(config.EXPERIMENTS_PATH))
        report = train_vs_test(candidates, train, reference, noisy)
        report.to_csv(MODEL_REPORT_PATH, index=False)
    summary = summarise(report)
    summary.to_csv(MODEL_SUMMARY_PATH, index=False)

    # Keep a copy of the scorer's chart under reports/ so the repository shows it.
    if SCORER_CHART.is_file():
        COMMITTED_CHART.write_bytes(SCORER_CHART.read_bytes())

    blocks = build_blocks(summary)
    render_docx(blocks, DOCX_PATH)
    render_markdown(blocks, MARKDOWN_PATH)
    log.info("wrote %s, %s, %s and %s", MODEL_REPORT_PATH.name, MODEL_SUMMARY_PATH.name,
             DOCX_PATH.name, MARKDOWN_PATH.name)
    print(summary.round(3).to_string(index=False))


if __name__ == "__main__":
    main()
